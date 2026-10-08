"""
DOCX Parser Module using python-docx with built-in zipfile/XML fallback.
Extracts text from paragraphs and tables in DOCX files with robust error handling.
"""

import io
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Union, BinaryIO
from modules.document_processing.file_validator import ValidationError
from modules.utils.logging_utils import get_logger

logger = get_logger(__name__)


class DocxParser:
    """Extracts raw text content from DOCX files using python-docx or standard XML parser."""

    @staticmethod
    def extract_text(file_obj: Union[str, Path, BinaryIO, bytes, io.BytesIO]) -> str:
        """
        Extracts text from DOCX file including paragraphs and tables.

        Args:
            file_obj: File path, stream, or bytes.

        Returns:
            Extracted text content as a single string.

        Raises:
            ValidationError: If DOCX is invalid or empty.
        """
        # 1. Try python-docx
        try:
            import docx
            if isinstance(file_obj, (bytes, bytearray)):
                stream = io.BytesIO(file_obj)
                doc = docx.Document(stream)
            elif isinstance(file_obj, (str, Path)):
                doc = docx.Document(str(file_obj))
            elif hasattr(file_obj, "read"):
                content = file_obj.read()
                if hasattr(file_obj, "seek"):
                    file_obj.seek(0)
                stream = io.BytesIO(content)
                doc = docx.Document(stream)
            else:
                doc = None

            if doc:
                text_parts = []
                for p in doc.paragraphs:
                    cleaned_p = p.text.strip()
                    if cleaned_p:
                        text_parts.append(cleaned_p)
                for table in doc.tables:
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            text_parts.append(" | ".join(row_cells))
                full_text = "\n".join(text_parts).strip()
                if full_text:
                    return full_text
        except ImportError:
            pass
        except Exception as e:
            logger.warning("python-docx extraction failed: %s, attempting standard XML fallback.", e)

        # 2. Built-in DOCX fallback using standard library zipfile and XML parsing
        try:
            if isinstance(file_obj, (str, Path)):
                archive = zipfile.ZipFile(str(file_obj))
            elif isinstance(file_obj, (bytes, bytearray)):
                archive = zipfile.ZipFile(io.BytesIO(file_obj))
            elif hasattr(file_obj, "read"):
                content = file_obj.read()
                if hasattr(file_obj, "seek"):
                    file_obj.seek(0)
                archive = zipfile.ZipFile(io.BytesIO(content))
            else:
                raise ValidationError("Invalid DOCX input.")

            with archive.open("word/document.xml") as xml_file:
                xml_content = xml_file.read()
                root = ET.fromstring(xml_content)
                # w:t represents text elements in OpenXML
                namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paragraphs = []
                for p_elem in root.iterfind(".//w:p", namespaces):
                    texts = [t.text for t in p_elem.iterfind(".//w:t", namespaces) if t.text]
                    if texts:
                        paragraphs.append("".join(texts).strip())

                full_text = "\n".join(paragraphs).strip()
                if full_text:
                    return full_text
        except Exception as e:
            logger.error("Standard XML DOCX extraction failed: %s", str(e))
            raise ValidationError(f"Error reading DOCX file: {str(e)}")

        raise ValidationError("No extractable text found in DOCX file.")
